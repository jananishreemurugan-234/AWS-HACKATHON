"""Unit tests for ResumePilot AI document parser service."""

import io
import unittest
from unittest.mock import MagicMock, patch
from types import SimpleNamespace
import docx
from PIL import Image
import pypdf

from services.resume_parser import (
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_text_with_tesseract,
    extract_text_with_textract,
    parse_resume_document,
    pdf_has_images,
)
import services.parser_service as parser_service_compat


def create_text_pdf(text: str) -> bytes:
    """Generate a clean, valid PDF containing the specified text string."""
    objs = [
        b"1 0 obj\n<</Type /Catalog /Pages 2 0 R>>\nendobj\n",
        b"2 0 obj\n<</Type /Pages /Kids [3 0 R] /Count 1>>\nendobj\n",
        b"3 0 obj\n<</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources <</Font <</F1 5 0 R>>>>>>\nendobj\n",
        f"4 0 obj\n<</Length {len(text) + 35}>>\nstream\nBT\n/F1 12 Tf\n72 712 Td\n({text}) Tj\nET\nendstream\nendobj\n".encode("latin1"),
        b"5 0 obj\n<</Type /Font /Subtype /Type1 /BaseFont /Helvetica>>\nendobj\n",
    ]
    header = b"%PDF-1.4\n"
    offsets = [0]
    curr = len(header)
    for obj in objs:
        offsets.append(curr)
        curr += len(obj)

    xref = f"xref\n0 {len(offsets)}\n".encode("ascii")
    xref += f"{offsets[0]:010d} 65535 f \r\n".encode("ascii")
    for off in offsets[1:]:
        xref += f"{off:010d} 00000 n \r\n".encode("ascii")

    trailer = f"trailer\n<</Size {len(offsets)} /Root 1 0 R>>\nstartxref\n{curr}\n%%EOF\n".encode("ascii")
    return header + b"".join(objs) + xref + trailer


def create_image_only_pdf() -> bytes:
    """Generate a valid single-page PDF containing only an embedded raster image."""
    img = Image.new("RGB", (150, 150), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PDF")
    return buf.getvalue()


def create_blank_pdf() -> bytes:
    """Generate a valid single-page PDF that has no text and no images."""
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def create_encrypted_pdf(password: str = "securepass") -> bytes:
    """Generate an encrypted PDF with user password protection."""
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


class TestParserService(unittest.TestCase):
    def test_parse_none_file(self):
        success, text, msg = parse_resume_document(None)
        self.assertFalse(success)
        self.assertEqual(text, "")
        self.assertIn("No file provided", msg)

    def test_parse_unsupported_extension(self):
        fake_file = SimpleNamespace(
            name="resume.txt",
            getvalue=lambda: b"plain text",
            size=10,
        )
        success, text, msg = parse_resume_document(fake_file)
        self.assertFalse(success)
        self.assertEqual(text, "")
        self.assertIn("Unsupported file format", msg)

    def test_parse_empty_file_pdf(self):
        fake_file = SimpleNamespace(
            name="empty_resume.pdf",
            getvalue=lambda: b"",
            size=0,
        )
        success, text, msg = parse_resume_document(fake_file)
        self.assertFalse(success)
        self.assertEqual(text, "")
        self.assertIn("empty", msg.lower())

    def test_parse_empty_file_docx(self):
        fake_file = SimpleNamespace(
            name="empty_resume.docx",
            getvalue=lambda: b"",
            size=0,
        )
        success, text, msg = parse_resume_document(fake_file)
        self.assertFalse(success)
        self.assertEqual(text, "")
        self.assertIn("empty", msg.lower())

    def test_parse_docx_content(self):
        doc = docx.Document()
        doc.add_paragraph("Jane Doe - Senior Python Engineer")
        doc.add_paragraph("Experienced in cloud architecture and microservices.")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        fake_file = SimpleNamespace(
            name="Jane_Doe_Resume.docx",
            getvalue=lambda: buffer.getvalue(),
            size=len(buffer.getvalue()),
        )

        success, text, msg = parse_resume_document(fake_file)
        self.assertTrue(success)
        self.assertIn("Jane Doe", text)
        self.assertIn("Senior Python Engineer", text)
        self.assertIn("Successfully extracted", msg)

    def test_parse_docx_empty_document(self):
        doc = docx.Document()
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        fake_file = SimpleNamespace(
            name="Empty_Content.docx",
            getvalue=lambda: buffer.getvalue(),
            size=len(buffer.getvalue()),
        )

        success, text, msg = parse_resume_document(fake_file)
        self.assertFalse(success)
        self.assertEqual(text, "")
        self.assertIn("no readable text was found", msg.lower())

    def test_parse_pdf_text_content(self):
        pdf_bytes = create_text_pdf("John Doe - Senior Software Engineer and DevOps Specialist")
        fake_file = SimpleNamespace(
            name="John_Doe_Resume.pdf",
            getvalue=lambda: pdf_bytes,
            size=len(pdf_bytes),
        )

        success, text, msg = parse_resume_document(fake_file)
        self.assertTrue(success)
        self.assertIn("John Doe", text)
        self.assertIn("Senior Software Engineer", text)
        self.assertIn("Successfully extracted", msg)

    def test_extract_text_from_pdf_direct(self):
        pdf_bytes = create_text_pdf("Direct PDF Text Stream Test")
        stream = io.BytesIO(pdf_bytes)
        extracted = extract_text_from_pdf(stream)
        self.assertIn("Direct PDF Text Stream Test", extracted)

    def test_parse_pdf_blank_page(self):
        blank_bytes = create_blank_pdf()
        fake_file = SimpleNamespace(
            name="Blank_Resume.pdf",
            getvalue=lambda: blank_bytes,
            size=len(blank_bytes),
        )

        success, text, msg = parse_resume_document(fake_file)
        self.assertFalse(success)
        self.assertEqual(text, "")
        self.assertIn("blank", msg.lower())

    @patch("services.resume_parser.extract_text_with_tesseract")
    def test_parse_pdf_scanned_image_only_with_ocr(self, mock_ocr):
        """Scanned PDF automatically uses local Tesseract OCR fallback when text extraction is empty."""
        mock_ocr.return_value = (
            True,
            "Jane Doe - Cloud Engineer",
            "Successfully extracted 25 characters from Scanned_Resume.pdf using local Tesseract OCR.",
        )

        image_pdf_bytes = create_image_only_pdf()
        fake_file = SimpleNamespace(
            name="Scanned_Resume.pdf",
            getvalue=lambda: image_pdf_bytes,
            size=len(image_pdf_bytes),
        )

        success, text, msg = parse_resume_document(fake_file)
        self.assertTrue(success)
        self.assertIn("Jane Doe", text)
        self.assertIn("local Tesseract OCR", msg)

    @patch("services.resume_parser.extract_text_with_tesseract")
    def test_parse_pdf_scanned_image_only_ocr_failure(self, mock_ocr):
        """When local OCR fails or Tesseract is missing, returns clear error message without fabricating text."""
        mock_ocr.return_value = (
            False,
            "",
            "Local Tesseract OCR is not installed or not found on your system PATH.",
        )

        image_pdf_bytes = create_image_only_pdf()
        fake_file = SimpleNamespace(
            name="Scanned_Resume.pdf",
            getvalue=lambda: image_pdf_bytes,
            size=len(image_pdf_bytes),
        )

        success, text, msg = parse_resume_document(fake_file)
        self.assertFalse(success)
        self.assertEqual(text, "")
        self.assertIn("tesseract", msg.lower())

    def test_parse_pdf_corrupted(self):
        fake_file = SimpleNamespace(
            name="corrupted.pdf",
            getvalue=lambda: b"invalid binary header %%PDF--corrupt-stream",
            size=40,
        )
        success, text, msg = parse_resume_document(fake_file)
        self.assertFalse(success)
        self.assertEqual(text, "")
        self.assertIn("corrupt", msg.lower())

    def test_parse_pdf_encrypted(self):
        enc_bytes = create_encrypted_pdf("secretpass")
        fake_file = SimpleNamespace(
            name="protected.pdf",
            getvalue=lambda: enc_bytes,
            size=len(enc_bytes),
        )
        success, text, msg = parse_resume_document(fake_file)
        self.assertFalse(success)
        self.assertEqual(text, "")
        self.assertTrue("encrypted" in msg.lower() or "password" in msg.lower())

    def test_parser_service_backwards_compatibility(self):
        self.assertEqual(
            parser_service_compat.parse_resume_document, parse_resume_document
        )
        self.assertEqual(
            parser_service_compat.extract_text_from_pdf, extract_text_from_pdf
        )
        self.assertEqual(
            parser_service_compat.extract_text_from_docx, extract_text_from_docx
        )
        self.assertEqual(
            parser_service_compat.extract_text_with_tesseract, extract_text_with_tesseract
        )
        self.assertEqual(
            parser_service_compat.extract_text_with_textract, extract_text_with_textract
        )
        self.assertEqual(parser_service_compat.pdf_has_images, pdf_has_images)


if __name__ == "__main__":
    unittest.main()
